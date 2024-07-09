import os
import shutil
import PyPDF2
from docx import Document
import xlrd
import openpyxl
import requests
import logging
import zipfile
import hashlib
import uuid

# Define your OpenAI API key here
OPENAI_API_KEY = "OPEN AI KEY HERE"

# Specify the directory containing the files to organize
directory = "C:\\Users\\DwainBrowne\\OneDrive - comtech.repair\\Documents"

# Define the category folders to skip
category_folders = ["_BILLING", "_FINANCIALS", "_MANUALS", "_NOTES", "_PERSONAL", "_PROJECTS", "_PHOTOS"]

# Setup logging
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

# Function to calculate the hash of a file for duplicate detection
def calculate_file_hash(file_path):
    hasher = hashlib.md5()
    with open(file_path, 'rb') as file:
        buf = file.read()
        hasher.update(buf)
    return hasher.hexdigest()

# Function to extract text from PDF files
def extract_text_from_pdf(file_path):
    logging.info(f"Extracting text from PDF: {file_path}")
    try:
        with open(file_path, 'rb') as file:
            reader = PyPDF2.PdfReader(file)
            text = ""
            for page in reader.pages:
                text += page.extract_text()
            return text[:500]  # Return first 500 characters
    except Exception as e:
        logging.error(f"Error reading PDF {file_path}: {e}")
        return ""

# Function to extract text from DOCX files
def extract_text_from_docx(file_path):
    logging.info(f"Extracting text from DOCX: {file_path}")
    try:
        doc = Document(file_path)
        text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
        return text[:500]  # Return first 500 characters
    except Exception as e:
        logging.error(f"Error reading DOCX {file_path}: {e}")
        return ""

# Function to extract text from XLS files
def extract_text_from_xls(file_path):
    logging.info(f"Extracting text from XLS: {file_path}")
    try:
        book = xlrd.open_workbook(file_path)
        sheet = book.sheet_by_index(0)
        text = "\n".join(["\t".join(map(str, sheet.row_values(row))) for row in range(min(sheet.nrows, 20))])
        return text[:500]  # Return first 500 characters
    except Exception as e:
        logging.error(f"Error reading XLS {file_path}: {e}")
        return ""

# Function to extract text from XLSX files
def extract_text_from_xlsx(file_path):
    logging.info(f"Extracting text from XLSX: {file_path}")
    try:
        book = openpyxl.load_workbook(file_path)
        sheet = book.active
        text = "\n".join(["\t".join(map(str, [cell.value for cell in row])) for row in sheet.iter_rows(min_row=1, max_row=20)])
        return text[:500]  # Return first 500 characters
    except Exception as e:
        logging.error(f"Error reading XLSX {file_path}: {e}")
        return ""

# Function to classify text using GPT
def classify_text(text, folder_name):
    logging.info("Classifying text with GPT")
    try:
        response = requests.post(
            "https://api.openai.com/v1/chat/completions",
            headers={
                "Authorization": f"Bearer {OPENAI_API_KEY}",
                "Content-Type": "application/json",
            },
            json={
                "model": "gpt-4o",
                "messages": [
                    {"role": "system", "content": "You are a file categorization assistant."},
                    {"role": "user", "content": f"Classify the following text into one of the categories: _BILLING, _FINANCIALS, _MANUALS, _NOTES, _PERSONAL, _PROJECTS or UNKNOWN.\n\n{text}\n\nFolder name hint: {folder_name}"}
                ],
                "max_tokens": 10,
            },
        )
        response.raise_for_status()  # Raise an error for bad HTTP status codes
        response_data = response.json()
        logging.info(f"API response: {response_data}")
        category = response_data["choices"][0]["message"]["content"].strip()
        logging.info(f"Text classified as: {category}")
        return category if category in ["_BILLING", "_FINANCIALS", "_MANUALS", "_NOTES", "_PERSONAL", "_PROJECTS"] else "UNKNOWN"
    except requests.exceptions.HTTPError as http_err:
        logging.error(f"HTTP error occurred: {http_err}")
        logging.error(f"API response content: {response.content}")
        return "UNKNOWN"
    except Exception as e:
        logging.error(f"Error classifying text: {e}")
        logging.error(f"API response content: {response.content if 'response' in locals() else 'No response content'}")
        return "UNKNOWN"

# Function to suggest a new file name based on text content
def suggest_file_name(text, folder_name, original_extension):
    logging.info("Suggesting file name with GPT")
    folder_name_sanitized = folder_name.replace("_", "-")
    try:
        response = requests.post(
            "https://api.openai.com/v1/chat/completions",
            headers={
                "Authorization": f"Bearer {OPENAI_API_KEY}",
                "Content-Type": "application/json",
            },
            json={
                "model": "gpt-4o",
                "messages": [
                    {"role": "system", "content": "You are a file renaming assistant."},
                    {"role": "user", "content": f"Based on the following text, suggest a concise and human-readable file name in the format ModelNumber-Year. DO NOT include the words Equipment, read the text and figure out a proper file name. Ensure the folder name is included as a hint: {folder_name_sanitized}\n\n{text}"}
                ],
                "max_tokens": 20,
            },
        )
        response.raise_for_status()  # Raise an error for bad HTTP status codes
        response_data = response.json()
        logging.info(f"API response: {response_data}")
        suggested_name = response_data["choices"][0]["message"]["content"].strip()
        logging.info(f"Suggested file name: {suggested_name}")

        # Sanitize the suggested name
        suggested_name = ''.join(c if c.isalnum() or c in ('-', '_') else '_' for c in suggested_name)
        # Limit the length of the suggested name
        suggested_name = suggested_name[:50]

        return f"{folder_name_sanitized}-{suggested_name}{original_extension}"
    except requests.exceptions.HTTPError as http_err:
        logging.error(f"HTTP error occurred: {http_err}")
        logging.error(f"API response content: {response.content}")
        return f"{folder_name_sanitized}-Unknown_File-{uuid.uuid4()}{original_extension}"
    except Exception as e:
        logging.error(f"Error suggesting file name: {e}")
        logging.error(f"API response content: {response.content if 'response' in locals() else 'No response content'}")
        return f"{folder_name_sanitized}-Unknown_File-{uuid.uuid4()}{original_extension}"

# Function to delete empty folders
def delete_empty_folders(path):
    if not os.path.isdir(path):
        return

    for root, dirs, files in os.walk(path, topdown=False):
        for dir in dirs:
            dir_path = os.path.join(root, dir)
            try:
                os.rmdir(dir_path)
                logging.info(f"Deleted empty folder: {dir_path}")
            except OSError:
                pass

# Function to organize files into categories
def organize_files(directory):
    processed_files = {}  # Dictionary to store file hashes for duplicate detection

    for root, _, files in os.walk(directory):
        # Skip category folders
        if os.path.basename(root) in category_folders:
            continue

        folder_name = os.path.basename(root)
        if "download" in folder_name.lower():
            folder_name = ""

        for file in files:
            file_path = os.path.join(root, file)
            file_extension = os.path.splitext(file)[1].lower()
            text = ""

            # Ensure the file still exists before processing
            if not os.path.exists(file_path):
                logging.warning(f"File not found: {file_path}")
                continue

            # Extract text based on file extension
            logging.info(f"Processing file: {file}")
            if file_extension in ['.jpg', '.jpeg', '.png']:
                # Move image files to _PHOTOS folder
                target_dir = os.path.join(directory, "_PHOTOS")
                os.makedirs(target_dir, exist_ok=True)
                file_hash = calculate_file_hash(file_path)
                if file_hash in processed_files:
                    logging.info(f"Duplicate photo detected: {file}. Skipping...")
                    continue
                new_file_path = os.path.join(target_dir, file)
                logging.info(f"Moving photo to {target_dir}")
                shutil.move(file_path, new_file_path)
                processed_files[file_hash] = new_file_path
                continue
            elif file_extension == '.pdf':
                text = extract_text_from_pdf(file_path)
            elif file_extension == '.docx':
                text = extract_text_from_docx(file_path)
            elif file_extension == '.xls':
                text = extract_text_from_xls(file_path)
            elif file_extension == '.xlsx':
                text = extract_text_from_xlsx(file_path)
            elif file_extension == '.zip':
                logging.info(f"Extracting ZIP file: {file_path}")
                try:
                    with zipfile.ZipFile(file_path, 'r') as zip_ref:
                        extract_dir = os.path.join(root, os.path.splitext(file)[0])
                        zip_ref.extractall(extract_dir)
                        organize_files(extract_dir)  # Recursively organize extracted files
                    # Remove the zip file after successful extraction
                    logging.info(f"Removing ZIP file: {file_path}")
                    os.remove(file_path)
                except zipfile.BadZipFile as e:
                    logging.error(f"Error extracting ZIP file {file_path}: {e}")
                continue
            else:
                logging.info(f"Skipping unknown file type: {file}")
                continue  # Skip unknown file types

            if text:
                # Classify the text to determine the category
                category = classify_text(text, folder_name)
                # Suggest a new file name based on the content and folder name
                new_file_name = suggest_file_name(text, folder_name, file_extension)
                # Determine the target directory based on the category
                target_dir = os.path.join(directory, category)
                os.makedirs(target_dir, exist_ok=True)
                
                # Calculate the hash of the file for duplicate detection
                file_hash = calculate_file_hash(file_path)
                if file_hash in processed_files:
                    logging.info(f"Duplicate file detected: {file}. Skipping...")
                    continue

                # Ensure the target directory exists
                if not os.path.exists(target_dir):
                    os.makedirs(target_dir, exist_ok=True)

                # Move and rename the file to the appropriate directory
                try:
                    logging.info(f"Moving file to {target_dir} with new name {new_file_name}")
                    shutil.move(file_path, os.path.join(target_dir, new_file_name))
                    processed_files[file_hash] = os.path.join(target_dir, new_file_name)
                except FileNotFoundError as e:
                    logging.error(f"Error moving file {file_path}: {e}")

    # Delete empty folders after processing
    delete_empty_folders(directory)

# Run the organization function
organize_files(directory)
